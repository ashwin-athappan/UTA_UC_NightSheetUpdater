import os
import base64
from collections import defaultdict
from datetime import datetime, timedelta

import environ
import requests
import openpyxl

from old.office365_api import Sharepoint

# === ENVIRONMENT SETUP ===
env = environ.Env()
env.read_env()

API_KEY = env("MAZEVO_API_KEY")
GET_EVENTS_URL = env("GET_EVENTS_URL")
GET_BOOKING_DETAILS_URL = env("GET_BOOKING_DETAILS_URL")
GET_DIAGRAM_URL = env("GET_DIAGRAM_URL")

IMAGES_DOWNLOADED_FLAG = False

BUILDING_IDS = [334, 335, 340, 341]
SHAREPOINT_URL_BASE = "https://mavsuta.sharepoint.com/sites/EOTEventOperationsTeam181"

sharepoint = Sharepoint()


def format_date(dt: datetime) -> str:
    return dt.strftime("%Y-%m-%dT00:00:00-05:00")


def get_date_str(date_str: str) -> str:
    dt = datetime.fromisoformat(date_str)
    return dt.strftime("%m_%d_%Y")


def fetch_api_data(API_URL: str, body, method='POST') -> dict:
    headers = {
        "X-API-Key": API_KEY,
        "Content-Type": "application/json"
    }
    try:
        if method.upper() == 'GET':
            response = requests.get(API_URL, headers=headers)
        else:
            response = requests.post(API_URL, json=body, headers=headers)
        response.raise_for_status()
        print(f"✅ API call to {API_URL} successful")
        return response.json()
    except Exception as e:
        raise Exception(f"API call failed: {e}")


def save_diagram_and_upload(base64_string: str, file_name: str, upload_folder: str, content_type: str) -> str:
    try:
        if not file_name.lower().endswith(".png"):
            file_name = file_name.rsplit('.', 1)[0] + "." + content_type.split("/")[1]

        binary_data = base64.b64decode(base64_string)

        local_path = f"./local_directory/temp_diagrams/{file_name}"
        os.makedirs(os.path.dirname(local_path), exist_ok=True)
        with open(local_path, "wb") as f:
            f.write(binary_data)

        result = sharepoint.upload_file(file_name, upload_folder, binary_data)
        if result["error"]:
            raise Exception(f"Upload failed: {result['error']}")

        return f"{SHAREPOINT_URL_BASE}/Shared%20Documents/{upload_folder}/{file_name}"
    except Exception as e:
        raise Exception(f"Error handling diagram upload: {e}")


def group_bookings_by_date(bookings: list, sheet_type: str) -> dict:
    grouped = defaultdict(list)
    for booking in bookings:
        try:
            dt = datetime.fromisoformat(booking["dateTimeStart"])
            if sheet_type == 'night_sheet':
                dt = dt - timedelta(days=1)
            key = dt.strftime("%m_%d_%Y")
            grouped[key].append(booking)
        except Exception as e:
            print(f"⚠️ Skipping booking with bad date: {e}")
    return grouped


def download_and_add_diagram_path(grouped_bookings: dict):
    image_folder_base = "General/EventSetupDiagrams/Mazevo/RoomDiagrams"
    image_folder_name = datetime.now().strftime("%Y_%m_%d")
    for sheet_name, bookings in grouped_bookings.items():
        folder_path = f"{image_folder_base}/{image_folder_name}"
        if not sharepoint.check_if_folder_exists(folder_path)['exists']:
            sharepoint.create_folder(folder_path)

        for booking in bookings:
            if booking.get("hasDiagram"):
                try:
                    diagram_url = f"{GET_DIAGRAM_URL}{booking['bookingId']}"
                    response = fetch_api_data(diagram_url, {}, method='GET')
                    base64_str = response.get("file")
                    content_type = response.get("contentType")
                    if base64_str:
                        diagram_file_name = f"{sheet_name}_{booking['bookingId']}_{response['fileName']}"
                        uploaded_path = save_diagram_and_upload(base64_str, diagram_file_name, folder_path, content_type)
                        booking["diagramPath"] = uploaded_path
                    else:
                        booking["diagramPath"] = None
                except Exception as e:
                    print(f"⚠️ Diagram error for booking {booking.get('bookingId')}: {e}")
                    booking["diagramPath"] = None
            else:
                booking["diagramPath"] = None
    return grouped_bookings


def write_bookings_to_excel(bookings_by_date: dict, file_path: str):
    wb = openpyxl.load_workbook(file_path)
    remaining_bookings = []

    for sheet_name, bookings in bookings_by_date.items():
        if sheet_name not in wb.sheetnames:
            print(f"❌ Sheet {sheet_name} not found.")
            continue

        sheet = wb[sheet_name]
        table = sheet.tables[list(sheet.tables)[0]]
        data_range = sheet[table.ref]
        headers = [cell.value for cell in data_range[0]]

        try:
            room_idx = headers.index("ROOM")
            end_idx = headers.index("END ")
            start_idx = headers.index("START")
            setup_idx = headers.index("SETUP")
            tech_idx = headers.index("TECH")
            notes_idx = headers.index("NOTES")
            drawings_idx = headers.index("DRAWINGS")
        except ValueError as e:
            print(f"❌ Column missing in {sheet_name}: {e}")
            continue

        rows = data_range[1:]
        inserted_rooms = []

        for i, row in enumerate(rows):
            row_room = row[room_idx].value
            for b in bookings:
                if b["roomDescription"] == row_room:
                    if b["roomDescription"] in inserted_rooms:
                        remaining_bookings.append(b)
                        continue

                    start_time = datetime.fromisoformat(b["dateTimeStart"]).strftime("%I:%M %p")
                    end_time = datetime.fromisoformat(b["dateTimeEnd"]).strftime("%I:%M %p")

                    row[start_idx].value = start_time
                    row[end_idx].value = end_time
                    row[notes_idx].value = b.get("setupNotes", "")

                    setup = f"{b['setupStyle']} for {b['setupCount']}" if b["setupStyle"] else ("See Notes" if not b["hasDiagram"] else "")
                    row[setup_idx].value = setup

                    tech_lines = []
                    for d in b.get("bookingDetails", []):
                        line = f"{d['resource']} - ({d['quantity']})"
                        if d.get("notes"):
                            line += f" - [{d['notes']}]"
                        tech_lines.append(line)
                    row[tech_idx].value = "\n".join(tech_lines)

                    if b["hasDiagram"] and b["diagramPath"]:
                        row[drawings_idx].hyperlink = b["diagramPath"]
                        row[drawings_idx].value = b["diagramPath"]

                    inserted_rooms.append(b["roomDescription"])

    wb.save(file_path)
    print("\nRemaining bookings not added due to duplicate rooms:")
    for b in remaining_bookings:
        print(f" - {b['roomDescription']} on {b['dateTimeStart']}")
    return remaining_bookings


def filter_events(api_data):
    return list({
        item["bookingId"]
        for item in api_data
        if item.get("statusDescription") == "Confirmed" and item.get("eventType") != "Maintenance"
    })


def process_excel_night_sheet(bookings: list, file_path: str):
    grouped = group_bookings_by_date(bookings, sheet_type='night_sheet')
    print("📅 Night Sheet Grouped bookings by date:")
    for date, items in grouped.items():
        print(f" - {date}: {len(items)} bookings")
    if not IMAGES_DOWNLOADED_FLAG:
        grouped = download_and_add_diagram_path(grouped)
    remaining_bookings = write_bookings_to_excel(grouped, file_path)
    return remaining_bookings

def process_excel_turnovers_sheet(bookings: list, file_path: str):
    grouped = group_bookings_by_date(bookings, sheet_type='turnovers')
    print("📅 Turnovers Grouped bookings by date:")
    for date, items in grouped.items():
        print(f" - {date}: {len(items)} bookings")
    write_bookings_to_excel(grouped, file_path)

def run_on_sharepoint_file(start_date: datetime, end_date: datetime, folder_path: str, night_sheet_filename: str = "Night Sheet - Multi Day Test.xlsx", turnovers_sheet_filename: str = "Turnovers - Multi Day Test.xlsx") -> str:
    events_body = {
        "start": format_date(start_date),
        "end": format_date(end_date),
        "buildingIds": BUILDING_IDS
    }

    events_data = fetch_api_data(GET_EVENTS_URL, events_body)
    booking_ids = filter_events(events_data)

    booking_details_body = {
        "bookingIds": booking_ids
    }

    booking_data = fetch_api_data(GET_BOOKING_DETAILS_URL, booking_details_body)

    result = sharepoint.download_file(night_sheet_filename, folder_path)
    if result["error"]:
        raise Exception(f"Download failed: {result['error']}")
    local_path = result["downloaded_file_path"]

    remaining_bookings = process_excel_night_sheet(booking_data, local_path)
    if remaining_bookings:
        result = sharepoint.download_file(turnovers_sheet_filename, folder_path)
        if result["error"]:
            raise Exception(f"Download failed: {result['error']}")
        local_path = result["downloaded_file_path"]
        process_excel_turnovers_sheet(remaining_bookings, local_path)

    with open(local_path, "rb") as f:
        content = f.read()
    upload_result = sharepoint.upload_file(night_sheet_filename, folder_path, content)
    if upload_result["error"]:
        raise Exception(f"Upload failed: {upload_result['error']}")

    return f"✅ Processed and uploaded '{night_sheet_filename}' to '{folder_path}'"


# Example execution
if __name__ == "__main__":
    folder_path = "Apps/Mazevo"
    night_sheet_file_name = "Night Sheet - Multi Day Test.xlsx"
    turnover_sheet_file_name = "Turnovers - Multi Day Test.xlsx"
    start_date = datetime(2025, 6, 24)
    end_date = datetime(2025, 6, 28)

    try:
        res = run_on_sharepoint_file(start_date, end_date, folder_path, night_sheet_file_name, turnover_sheet_file_name)
        print(res)
    except Exception as e:
        print(f"❌ Error: {e}")
